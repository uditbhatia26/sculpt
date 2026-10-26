"""
test_hardening.py
=================
DB-free, network-free regression tests (run with `python -m unittest discover -s tests`
or `pytest tests`). Covers: auth refresh flags, schema validators, rate limiter, DOCX
extraction, CORS origin policy, password reset, and the optimization workflow helpers.
"""

import asyncio
import io
import os
import sys
import unittest
import uuid
from datetime import datetime, timedelta
from unittest.mock import MagicMock, patch

BACKEND_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if BACKEND_DIR not in sys.path:
    sys.path.insert(0, BACKEND_DIR)

from fastapi import HTTPException
from fastapi.testclient import TestClient
from pydantic import ValidationError

import main
from main import app, get_current_user, get_db
from config.rate_limit import rate_limit, reset_rate_limits
from models import workflow
from schema.schema import CalculateATS, OptimizeResumeRequest, GeneratePDFRequest


def make_user(**over):
    u = MagicMock()
    u.id = uuid.uuid4()
    u.email = "someone@example.com"
    u.full_name = "Some One"
    u.plan = "free"
    u.resume_yaml = "name: x"
    u.resume_filename = "r.pdf"
    u.email_verified = True
    u.is_byok_allowed = True
    u.is_admin = True
    u.selected_template = "jakes_resume"
    u.password_reset_token_hash = None
    u.password_reset_expires_at = None
    for k, v in over.items():
        setattr(u, k, v)
    return u


def run(coro):
    return asyncio.new_event_loop().run_until_complete(coro)


class TestAuthRefresh(unittest.TestCase):
    def tearDown(self):
        app.dependency_overrides.clear()

    def test_refresh_keeps_byok_and_admin_flags(self):
        user = make_user()
        app.dependency_overrides[get_current_user] = lambda: user
        app.dependency_overrides[get_db] = lambda: MagicMock()
        with patch.object(main, "get_weekly_usage", return_value=2), \
             patch.object(main, "get_daily_usage", return_value=1), \
             patch.object(main, "get_monthly_usage", return_value=3):
            res = TestClient(app).post("/auth/refresh")
        self.assertEqual(res.status_code, 200, res.text)
        body = res.json()
        self.assertTrue(body["is_byok_allowed"])
        self.assertTrue(body["is_admin"])
        self.assertEqual(body["weekly_usage"], 2)


class TestSchemaValidators(unittest.TestCase):
    JD = "x" * 60

    def test_bad_jd_cache_id_rejected(self):
        with self.assertRaises(ValidationError):
            CalculateATS(job_desc=self.JD, jd_cache_id="not-a-uuid")
        with self.assertRaises(ValidationError):
            OptimizeResumeRequest(job_desc=self.JD, jd_cache_id="'; DROP TABLE users;--")

    def test_good_jd_cache_id_accepted(self):
        u = str(uuid.uuid4())
        self.assertEqual(CalculateATS(job_desc=self.JD, jd_cache_id=u).jd_cache_id, u)
        self.assertIsNone(CalculateATS(job_desc=self.JD).jd_cache_id)

    def test_template_validated(self):
        with self.assertRaises(ValidationError):
            GeneratePDFRequest(template="evil_template")
        with self.assertRaises(ValidationError):
            OptimizeResumeRequest(job_desc=self.JD, template="../../etc/passwd")
        self.assertEqual(GeneratePDFRequest(template="classic_ivy").template, "classic_ivy")
        # No template means "use the user's saved one"
        self.assertIsNone(GeneratePDFRequest().template)
        self.assertIsNone(OptimizeResumeRequest(job_desc=self.JD).template)


class TestRateLimiter(unittest.TestCase):
    def setUp(self):
        reset_rate_limits()

    def _req(self, ip="1.2.3.4"):
        r = MagicMock()
        r.headers = {}
        r.client.host = ip
        return r

    def test_blocks_after_limit_per_ip(self):
        dep = rate_limit("t", 3, 60)
        for _ in range(3):
            dep(self._req())
        with self.assertRaises(HTTPException) as ctx:
            dep(self._req())
        self.assertEqual(ctx.exception.status_code, 429)
        self.assertIn("Retry-After", ctx.exception.headers)
        dep(self._req("9.9.9.9"))  # other IP unaffected


class TestCors(unittest.TestCase):
    def _req(self, origin):
        r = MagicMock()
        r.headers = {"origin": origin} if origin else {}
        return r

    def test_origin_policy(self):
        self.assertEqual(main._cors_origin(self._req("chrome-extension://abc")), "chrome-extension://abc")
        self.assertEqual(main._cors_origin(self._req("http://localhost:5173")), "http://localhost:5173")
        self.assertEqual(main._cors_origin(self._req("https://evil.example")), "")
        self.assertEqual(main._cors_origin(self._req("")), "")


class TestDocxExtraction(unittest.TestCase):
    def test_extracts_paragraphs_and_tables(self):
        from docx import Document
        doc = Document()
        doc.add_paragraph("Jane Doe")
        doc.add_paragraph("Software Engineer")
        t = doc.add_table(rows=1, cols=2)
        t.rows[0].cells[0].text = "Python"
        t.rows[0].cells[1].text = "FastAPI"
        buf = io.BytesIO()
        doc.save(buf)
        text = main.extract_docx_text(buf.getvalue())
        self.assertIn("Jane Doe", text)
        self.assertIn("Python | FastAPI", text)


class TestPasswordReset(unittest.TestCase):
    def tearDown(self):
        app.dependency_overrides.clear()
        reset_rate_limits()

    def _client_with_user(self, user):
        db = MagicMock()
        db.query.return_value.filter.return_value.first.return_value = user
        app.dependency_overrides[get_db] = lambda: db
        return TestClient(app), db

    def test_forgot_password_does_not_leak_account_existence(self):
        client, _ = self._client_with_user(None)
        res = client.post("/auth/forgot-password", json={"email": "ghost@example.com"})
        self.assertEqual(res.status_code, 200)
        self.assertIn("If an account exists", res.json()["message"])

    def test_forgot_password_stores_hash_not_token(self):
        user = make_user()
        client, db = self._client_with_user(user)
        with patch.object(main, "send_password_reset_email") as send:
            res = client.post("/auth/forgot-password", json={"email": user.email})
        self.assertEqual(res.status_code, 200)
        token = send.call_args[0][2]
        self.assertNotEqual(user.password_reset_token_hash, token)
        self.assertEqual(user.password_reset_token_hash, main._hash_reset_token(token))

    def test_reset_with_valid_token(self):
        user = make_user(
            password_reset_token_hash=main._hash_reset_token("tok" * 5),
            password_reset_expires_at=datetime.utcnow() + timedelta(minutes=5),
        )
        client, _ = self._client_with_user(user)
        res = client.post("/auth/reset-password", json={"token": "tok" * 5, "new_password": "newpass1"})
        self.assertEqual(res.status_code, 200, res.text)
        self.assertIsNone(user.password_reset_token_hash)  # single use

    def test_reset_with_expired_or_unknown_token(self):
        expired = make_user(
            password_reset_token_hash="h",
            password_reset_expires_at=datetime.utcnow() - timedelta(minutes=1),
        )
        client, _ = self._client_with_user(expired)
        self.assertEqual(
            client.post("/auth/reset-password", json={"token": "t" * 12, "new_password": "newpass1"}).status_code, 400)
        client, _ = self._client_with_user(None)
        self.assertEqual(
            client.post("/auth/reset-password", json={"token": "t" * 12, "new_password": "newpass1"}).status_code, 400)

    def test_reset_page_sanitises_token(self):
        client, _ = self._client_with_user(None)
        res = client.get("/auth/reset-password", params={"token": "abc');alert(1);//"})
        self.assertEqual(res.status_code, 200)
        self.assertNotIn("alert(1)", res.text)


class TestPdfRoutesRequireVerifiedEmail(unittest.TestCase):
    def tearDown(self):
        app.dependency_overrides.clear()

    def test_generate_pdf_blocked_for_unverified(self):
        app.dependency_overrides[get_current_user] = lambda: make_user(email_verified=False)
        app.dependency_overrides[get_db] = lambda: MagicMock()
        res = TestClient(app).post("/generate-pdf", json={})
        self.assertEqual(res.status_code, 403)


# ── Workflow ────────────────────────────────────────────────────────────────

class TestWorkflow(unittest.TestCase):
    def test_best_candidate_prefers_page_fit(self):
        state = {
            "target_page_count": 1,
            "score_history": [
                {"score": 90, "yaml": "a", "page_count": 2},
                {"score": 80, "yaml": "b", "page_count": 1},
            ],
        }
        self.assertEqual(workflow._best_candidate(state)["yaml"], "b")

    def test_best_candidate_falls_back_to_highest_score(self):
        state = {
            "target_page_count": 1,
            "score_history": [
                {"score": 90, "yaml": "a", "page_count": 2},
                {"score": 70, "yaml": "b", "page_count": 3},
            ],
        }
        self.assertEqual(workflow._best_candidate(state)["yaml"], "a")
        self.assertIsNone(workflow._best_candidate({}))

    def test_page_budget_enforced_for_all_templates(self):
        state = {"selected_template": "jakes_resume", "target_page_count": 2, "current_resume_yaml": "name: x"}
        with patch.object(workflow, "render_latex_template", return_value="tex"), \
             patch.object(workflow, "compile_latex_to_pdf", return_value=b"pdf"), \
             patch.object(workflow, "PdfReader") as reader:
            reader.return_value.pages = [1, 2, 3]
            out = workflow.node_page_budget(state)
        self.assertFalse(out["page_budget_passed"])
        self.assertEqual(out["actual_page_count"], 3)
        self.assertTrue(out["critique_notes"])

    def test_compile_failure_is_surfaced_not_guessed(self):
        state = {"selected_template": "compact_onepage", "target_page_count": 1, "current_resume_yaml": "name: x"}
        with patch.object(workflow, "render_latex_template", side_effect=RuntimeError("boom")):
            out = workflow.node_page_budget(state)
        self.assertIsNone(out["actual_page_count"])
        self.assertEqual(out["compile_error"], "RuntimeError")
        self.assertTrue(out["page_budget_passed"])

    def test_draft_resets_critique_and_builds_on_best_candidate(self):
        captured = {}

        class FakeChain:
            async def ainvoke(self, input):
                captured.update(input)
                return MagicMock(content="name: improved")

        state = {
            "llm": object(),
            "selected_template": "jakes_resume",
            "original_resume_yaml": "name: orig",
            "current_resume_yaml": "name: regressed",
            "parsed_jd": {"job_title": "Eng", "skills": ["Python"], "job_description": "desc"},
            "critique_notes": ["fix this"],
            "target_page_count": 2,
            "score_history": [
                {"score": 60, "yaml": "name: best", "page_count": 1},
                {"score": 40, "yaml": "name: regressed", "page_count": 1},
            ],
        }
        with patch.object(workflow, "build_optimization_chain", return_value=FakeChain()):
            out = run(workflow.node_draft(state))
        self.assertEqual(captured["resume_yaml"], "name: best")
        self.assertIn("fix this", captured["template_guidance"])
        self.assertEqual(out["critique_notes"], [])

    def test_finalize_keywords_added_is_a_true_diff(self):
        state = {
            "selected_template": "jakes_resume",
            "target_page_count": 2,
            "original_resume_yaml": "skills:\n  - Python\n",
            "current_resume_yaml": "skills:\n  - Python\n  - Kubernetes\n",
            "parsed_jd": {"job_title": "Eng", "skills": ["Python", "Kubernetes", "Rust"], "job_description": "d"},
            "page_budget_passed": True,
            "ats_score_passed": True,
            "current_ats_score": 80.0,
            "original_ats_score": 60.0,
            "actual_page_count": 1,
            "rendered_latex": "tex",
            "current_ats_details": {"keyword_analysis": {"matched_keywords": ["Python", "Kubernetes"]}},
        }
        out = workflow.node_finalize(state)
        self.assertEqual(out["keywords_added"], ["Kubernetes"])
        self.assertTrue(any("Page budget met" in s for s in out["improvements_made"]))

    def test_router(self):
        self.assertEqual(workflow.should_continue({"page_budget_passed": True, "ats_score_passed": True}), "finalize")
        self.assertEqual(
            workflow.should_continue({"page_budget_passed": False, "ats_score_passed": True, "iteration_count": 0, "max_retries": 2}),
            "draft")
        self.assertEqual(
            workflow.should_continue({"page_budget_passed": False, "ats_score_passed": True, "iteration_count": 2, "max_retries": 2}),
            "finalize")


if __name__ == "__main__":
    unittest.main()
